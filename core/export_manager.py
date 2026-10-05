# ~/Desktop/acropdf/core/export_manager.py
import fitz
from core.file_io import publish_exclusive
import os
import tempfile

class ExportManager:
    def __init__(self, doc):
        self._doc = doc

    @property
    def _fitz(self) -> fitz.Document:
        return self._doc.fitz_doc

    def _safe_fitz(self) -> fitz.Document:
        doc = self._fitz
        if doc is None:
            raise RuntimeError("尚未載入文件")
        return doc

    def export(self, output_path: str, fmt: str, dpi: int = 150):
        self._safe_fitz()
        normalized = fmt.lower().replace("/a", "a")
        if normalized == "pdfa":
            return self.export_pdfa(output_path)
        if normalized in {"png", "jpg", "jpeg", "tiff"}:
            suffix = os.path.splitext(output_path)[1]
            output_dir = os.path.dirname(output_path) or "." if suffix else output_path
            output_dir = self._validated_output_dir(output_dir)
            base = os.path.splitext(os.path.basename(output_path))[0] if suffix else None
            return self.export_images(output_dir, fmt=normalized.replace("jpeg", "jpg"), dpi=dpi, base=base)
        exporters = {"docx": self.export_docx, "xlsx": self.export_xlsx, "pptx": self.export_pptx,
                     "txt": self.export_txt, "html": self.export_html, "pdf": self.export_pdf}
        if normalized not in exporters:
            raise ValueError(f"不支援的匯出格式：{fmt}")
        output_path = self._validated_output_path(output_path)
        source = self._doc.source_path
        if source and os.path.realpath(source) == os.path.realpath(output_path):
            raise ValueError("匯出請選擇新檔名，避免覆寫目前開啟的來源檔。")
        fd, temporary = tempfile.mkstemp(prefix="acropdf_export_", suffix="." + normalized,
                                         dir=os.path.dirname(os.path.abspath(output_path)))
        os.close(fd)
        try:
            exporters[normalized](temporary)
            with open(temporary, "r+b") as handle:
                os.fsync(handle.fileno())
            os.replace(temporary, output_path)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
        return output_path

    def _validated_output_path(self, output_path: str) -> str:
        output_path = os.path.expanduser(output_path or "")
        if not output_path:
            raise ValueError("請先選擇匯出檔案位置。")
        parent = os.path.dirname(output_path) or "."
        self._ensure_writable_dir(parent)
        return output_path

    def _validated_output_dir(self, output_dir: str) -> str:
        output_dir = os.path.expanduser(output_dir or "")
        if not output_dir:
            raise ValueError("請先選擇匯出資料夾。")
        self._ensure_writable_dir(output_dir)
        return output_dir

    @staticmethod
    def _ensure_writable_dir(path: str):
        if not os.path.isdir(path):
            raise FileNotFoundError(f"匯出資料夾不存在：{path}")
        if not os.access(path, os.W_OK):
            raise PermissionError(f"儲存位置不可寫入：{path}\n請改存到「文件」或「桌面」。")

    # ── Word ─────────────────────────────────────────────────────
    def export_docx(self, output_path: str):
        self._safe_fitz()  # raises RuntimeError if None
        from docx import Document
        from docx.shared import Pt, RGBColor
        word_doc = Document()
        page_count = self._fitz.page_count
        for i in range(page_count):
            page = self._fitz[i]
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if block["type"] == 0:  # 文字 block
                    for line in block["lines"]:
                        text = " ".join(span["text"] for span in line["spans"])
                        if text.strip():
                            p = word_doc.add_paragraph(text)
                            # 嘗試保留字型大小
                            if line["spans"]:
                                fs = line["spans"][0].get("size", 11)
                                for run in p.runs:
                                    run.font.size = Pt(fs)
            if i < page_count - 1:
                word_doc.add_page_break()
        word_doc.save(output_path)

    # ── Excel（表格偵測）─────────────────────────────────────────
    def export_xlsx(self, output_path: str):
        self._safe_fitz()
        import openpyxl
        wb = openpyxl.Workbook()
        for i in range(self._fitz.page_count):
            page = self._fitz[i]
            ws = wb.create_sheet(f"頁面{i+1}")
            # PyMuPDF 1.23+ 自動偵測表格
            found = False
            if hasattr(page, "find_tables"):
                tabs = page.find_tables()
                row_offset = 1
                for tab in tabs:
                    found = True
                    for r, row in enumerate(tab.extract()):
                        for c, cell in enumerate(row):
                            result = ws.cell(row=row_offset + r, column=c + 1, value=cell or "")
                            # PDF contents are data, never spreadsheet formulas.
                            if isinstance(cell, str):
                                result.data_type = "s"
                    row_offset += len(tab.extract()) + 2
            if not found:
                text = page.get_text()
                for r, line in enumerate(text.splitlines()):
                    cell = ws.cell(row=r + 1, column=1, value=line)
                    cell.data_type = "s"
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        wb.save(output_path)

    # ── PowerPoint ───────────────────────────────────────────────
    def export_pptx(self, output_path: str):
        from pptx import Presentation
        doc = self._safe_fitz()
        prs = Presentation()
        if doc.page_count:
            first = doc[0].rect
            prs.slide_width = int(first.width * 12700)
            prs.slide_height = int(first.height * 12700)
        for i in range(doc.page_count):
            page = doc[i]
            slide_layout = prs.slide_layouts[6]  # Blank
            slide = prs.slides.add_slide(slide_layout)
            mat = fitz.Matrix(2, 2)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            with tempfile.NamedTemporaryFile(prefix="acropdf_slide_", suffix=".png", delete=False) as tmp:
                img_path = tmp.name
            try:
                pix.save(img_path)
                ratio = min(prs.slide_width / pix.width, prs.slide_height / pix.height)
                width, height = int(pix.width * ratio), int(pix.height * ratio)
                slide.shapes.add_picture(img_path, left=(prs.slide_width - width) // 2,
                                         top=(prs.slide_height - height) // 2, width=width, height=height)
            finally:
                if os.path.exists(img_path):
                    os.remove(img_path)
        prs.save(output_path)

    # ── 圖片（每頁）─────────────────────────────────────────────
    def export_images(self, output_dir: str, fmt: str = "png", dpi: int = 150, base: str | None = None):
        doc = self._safe_fitz()
        if not 72 <= dpi <= 600:
            raise ValueError("圖片解析度須介於 72 與 600 DPI。")
        base = base or os.path.splitext(os.path.basename(self._doc.path or "page"))[0]
        paths = [os.path.join(output_dir, f"{base}_p{i+1:03d}.{fmt}") for i in range(doc.page_count)]
        for path in paths:
            if os.path.exists(path):
                raise FileExistsError(f"圖片已存在，請使用新檔名或資料夾：{path}")
        published = []
        with tempfile.TemporaryDirectory(prefix="acropdf_images_", dir=output_dir) as stage:
            staged = []
            for i, path in enumerate(paths):
                page = doc[i]
                pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
                temporary = os.path.join(stage, os.path.basename(path))
                if fmt == "tiff":
                    from PIL import Image
                    Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(temporary, format="TIFF")
                else:
                    pix.save(temporary)
                staged.append(temporary)
            try:
                for temporary, path in zip(staged, paths):
                    # Publication refuses a concurrently-created destination.
                    publish_exclusive(temporary, path)
                    published.append(path)
            except Exception:
                for path in published:
                    os.unlink(path)
                raise
        return paths

    # ── 純文字 ───────────────────────────────────────────────────
    def export_txt(self, output_path: str):
        self._safe_fitz()
        text = ""
        for i in range(self._fitz.page_count):
            text += self._fitz[i].get_text() + "\n\f\n"
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)

    # ── HTML ─────────────────────────────────────────────────────
    def export_html(self, output_path: str):
        self._safe_fitz()
        html_parts = ['<!DOCTYPE html><html><body>']
        for i in range(self._fitz.page_count):
            html_parts.append(self._fitz[i].get_text("html"))
            html_parts.append('<hr/>')
        html_parts.append('</body></html>')
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(html_parts))

    def export_pdf(self, output_path: str):
        self._safe_fitz().save(output_path, garbage=4, deflate=True)

    def export_pdfa(self, output_path: str):
        from core.pdf_standards import PDFStandards
        self.last_standard_report = PDFStandards(self._doc).export_pdfa(output_path)
        return output_path

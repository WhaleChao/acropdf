"""Commercial workflow regressions: data integrity, rendering and interaction state."""
import os
from unittest.mock import patch

import fitz
import pytest
from PyQt6.QtCore import Qt, QSettings, QPoint, QThreadPool
from PyQt6.QtGui import QCloseEvent, QImage, QPixmap, QPalette
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QMessageBox, QApplication

from app.config import Config
from core.document import PDFDocument
from ui.main_window import MainWindow
from ui.theme import COLORS, theme_manager
from ui.viewer.pdf_view import PDFView
from ui.viewer.page_widget import PageWidget
from core.ai_endpoint import endpoint_is_local, confirm_document_transfer


def opened(path):
    doc = PDFDocument()
    assert doc.open(str(path))
    return doc


@pytest.fixture
def window(qapp):
    w = MainWindow()
    w.show()
    yield w
    QThreadPool.globalInstance().waitForDone(10000)
    for d in w._docs:
        d._modified = False
    w.close()
    qapp.processEvents()


@pytest.fixture
def encrypted(tmp_path):
    path = tmp_path / 'locked.pdf'
    with fitz.open() as doc:
        doc.new_page().insert_text((72, 72), 'Encrypted contents')
        doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='user')
    return path


def test_new_document_is_saveable_and_dirty(tmp_path):
    doc = PDFDocument(); doc.new()
    assert doc.page_count == 1 and doc.is_modified
    assert doc.save(str(tmp_path / 'new.pdf'))
    assert not doc.is_modified
    doc.close()


def test_missing_save_path_does_not_write_current_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    doc = PDFDocument(); doc.new()
    assert not doc.save()
    assert not list(tmp_path.glob('acropdf_save_*'))
    doc.close()


def test_failed_atomic_save_preserves_original(sample_pdf, monkeypatch):
    doc = opened(sample_pdf)
    before = sample_pdf.read_bytes()
    doc.pages.rotate([0], 90)
    def fail(*args): raise OSError('disk full')
    monkeypatch.setattr('core.document.os.replace', fail)
    assert not doc.save()
    assert doc.is_modified and 'disk full' in doc.last_error
    assert sample_pdf.read_bytes() == before
    assert not list(sample_pdf.parent.glob('acropdf_save_*'))
    doc.close()


def test_save_preserves_destination_permissions(sample_pdf):
    sample_pdf.chmod(0o640)
    expected_mode = sample_pdf.stat().st_mode & 0o777
    doc = opened(sample_pdf); doc.pages.rotate([0], 90)
    assert doc.save()
    assert sample_pdf.stat().st_mode & 0o777 == expected_mode
    doc.close()


def test_office_source_never_overwritten(tmp_path):
    from docx import Document
    path = tmp_path / 'source.docx'
    word = Document(); word.add_paragraph('Original Office source'); word.save(path)
    original = path.read_bytes(); doc = opened(path)
    assert doc.path is None and doc.is_modified and doc.fitz_doc.is_pdf
    assert not doc.save(str(path))
    assert path.read_bytes() == original
    assert doc.save(str(tmp_path / 'converted.pdf'))
    assert path.read_bytes() == original
    doc.close()


def test_converted_image_is_an_editable_pdf(tmp_path):
    image = tmp_path / 'image.png'
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0,0,120,160), False); pix.clear_with(255); pix.save(image)
    doc = opened(image)
    assert doc.fitz_doc.is_pdf and doc.path is None
    doc.pages.insert_blank(1)
    assert doc.save(str(tmp_path / 'image.pdf'))
    assert doc.page_count == 2
    doc.close()


def test_encrypted_save_and_undo_preserve_password(encrypted):
    doc = PDFDocument(); assert doc.open(str(encrypted), 'user')
    doc.pages.rotate([0], 90)
    assert doc.undo() and doc.fitz_doc[0].rotation == 0
    assert doc.redo() and doc.fitz_doc[0].rotation == 90
    assert doc.save()
    with fitz.open(encrypted) as check:
        assert check.needs_pass and check.authenticate('user')
        assert check[0].rotation == 90
    doc.close()


@pytest.mark.parametrize('whole_window', [False, True])
def test_close_keeps_document_on_save_failure(window, sample_pdf, whole_window):
    window.open_file(str(sample_pdf)); doc = window._current_doc(); doc.pages.rotate([0],90)
    choice = QMessageBox.StandardButton.SaveAll if whole_window else QMessageBox.StandardButton.Save
    with patch.object(QMessageBox, 'question', return_value=choice), patch.object(QMessageBox, 'warning'), patch.object(doc, 'save', return_value=False):
        if whole_window:
            event = QCloseEvent(); window.closeEvent(event); assert not event.isAccepted()
        else:
            window._close_tab(0)
    assert window._doc_tabs.count() == 1 and window._current_doc() is doc
    assert doc.fitz_doc is not None


def test_cancel_save_of_new_document_keeps_tab(window):
    window.new_document()
    with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.Save), patch('ui.main_window.QFileDialog.getSaveFileName',return_value=('', '')):
        window._close_tab(0)
    assert window._doc_tabs.count() == 1


def test_close_new_document_adds_pdf_suffix(window, tmp_path):
    window.new_document()
    with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.Save), patch('ui.main_window.QFileDialog.getSaveFileName',return_value=(str(tmp_path/'new'), 'PDF 檔案 (*.pdf)')):
        window._close_tab(0)
    assert (tmp_path/'new.pdf').exists() and window._doc_tabs.count() == 0


def test_tab_reorder_keeps_document_and_view_in_sync(window, sample_pdf, tmp_path):
    other = tmp_path/'other.pdf'
    with fitz.open() as doc: doc.new_page(); doc.save(other)
    window.open_file(str(sample_pdf)); first = window._current_doc()
    window.open_file(str(other)); second = window._current_doc()
    window._detachable_bar.moveTab(1, 0)
    assert window._docs == [second, first]
    window._doc_tabs.setCurrentIndex(0)
    assert window._current_view()._doc is window._current_doc() is second
    window._doc_tabs.setCurrentIndex(1)
    assert window._current_view()._doc is first


def test_inactive_view_signals_do_not_change_current_controls(window, sample_pdf, tmp_path):
    window.open_file(str(sample_pdf)); old_view = window._current_view()
    path = tmp_path/'one.pdf'
    with fitz.open() as d: d.new_page(); d.save(path)
    window.open_file(str(path)); window._current_view().set_zoom(0.75)
    old_view.zoom_changed.emit(2.0); old_view.page_changed.emit(2)
    assert window._zoom_label.text() == '75%'
    assert window._page_spin.value() == 1


def test_split_view_has_two_loaded_views_and_keeps_position(window, sample_pdf, qapp):
    window.open_file(str(sample_pdf)); qapp.processEvents()
    window._current_view().set_zoom(.5); window._current_view().go_to_page(1)
    old = window._current_view()
    window._split_action.setChecked(True); window._toggle_split_view(); qapp.processEvents()
    split = window._doc_tabs.currentWidget()
    assert not split.secondary_view.isHidden()
    assert split.secondary_view._doc is window._current_doc()
    assert split.primary_view.current_page() == 1
    assert old._doc is None
    window._switch_theme('dark')
    for view in (split.primary_view, split.secondary_view):
        assert view.viewport().palette().color(QPalette.ColorRole.Window).name() == COLORS['dark']['canvas']
    window._split_action.setChecked(False); window._toggle_split_view()
    assert window._current_view().current_page() == 1


def test_closing_tab_disconnects_view(window, sample_pdf):
    window.open_file(str(sample_pdf)); view = window._current_view(); doc = window._current_doc()
    window._close_tab(0)
    assert view._doc is None and doc not in window._doc_connections


@pytest.mark.parametrize('mode',['light','dark','system'])
def test_theme_applies_to_application_dialogs_and_new_documents(window, mode, qapp):
    window._switch_theme(mode); window.new_document()
    manager=theme_manager()
    assert Config().theme == mode
    assert manager.mode == mode
    assert window._theme_combo.currentData() == mode
    assert window._theme_actions[mode].isChecked()
    assert qapp.palette().color(QPalette.ColorRole.Window).name() == COLORS[manager.resolved]['bg']
    assert window._current_view().viewport().palette().color(QPalette.ColorRole.Window).name() == COLORS[manager.resolved]['canvas']


def test_system_theme_reacts_and_explicit_theme_is_stable(qapp):
    manager=theme_manager(); manager.set_mode('system')
    with patch.object(manager,'apply') as apply:
        manager._system_changed(Qt.ColorScheme.Dark); apply.assert_called_once()
    manager.set_mode('light')
    with patch.object(manager,'apply') as apply:
        manager._system_changed(Qt.ColorScheme.Dark); apply.assert_not_called()


def test_document_actions_and_tool_search(window):
    assert not window._save_action.isEnabled() and window._mark_toolbar.isHidden()
    window._tools_panel._filter.setText('xyz-not-a-tool')
    assert not window._tools_panel._empty.isHidden()
    window._tools_panel._filter.setText('匯出')
    visible = [b.text() for b,_ in window._tools_panel._buttons if not b.isHidden()]
    assert visible and all('匯出' in name for name in visible)
    window.new_document()
    assert window._save_action.isEnabled() and not window._mark_toolbar.isHidden()


def test_command_palette_filters_and_executes_with_keyboard(window):
    from ui.widgets.command_palette import CommandPalette
    calls=[]
    palette=CommandPalette(window,[('開啟文件','open pdf',lambda:calls.append('open'),True),('儲存','save',lambda:calls.append('save'),False)])
    palette.show(); palette._input.setText('open')
    assert palette._results.count() == 1
    QTest.keyClick(palette._input, Qt.Key.Key_Return)
    assert calls == ['open']
    palette.deleteLater()


def test_clearing_search_resets_navigation(window, sample_pdf):
    window.open_file(str(sample_pdf)); window._search_bar.focus_input()
    window._search_bar._input.setText('Sample')
    assert len(window._search_results) == 3
    QTest.keyClick(window._search_bar._input, Qt.Key.Key_Return, Qt.KeyboardModifier.ShiftModifier)
    assert window._search_index == 2
    window._search_bar._input.clear()
    assert not window._search_results and window._search_index == -1
    assert not window._search_bar._next_btn.isEnabled()


def test_rotated_page_coordinates(qapp):
    with fitz.open() as doc:
        page=doc.new_page(width=200,height=300); page.set_rotation(90)
        widget=PageWidget(0); widget._zoom=2
        original=fitz.Point(30,40); visible=original * page.rotation_matrix
        recovered=widget.widget_to_pdf(QPoint(int(visible.x*2),int(visible.y*2)),page)
        assert abs(recovered.x-original.x)<.01 and abs(recovered.y-original.y)<.01


def test_rotation_rebuilds_geometry_and_discards_stale_render(qapp, sample_pdf):
    doc=opened(sample_pdf); view=PDFView(); view.load_document(doc)
    old_gen=view._render_gen; old_width=view._page_widgets[0].width()
    doc.pages.rotate([0],90)
    assert view._render_gen>old_gen and view._page_widgets[0].width()!=old_width
    old = QPixmap(5,5); view._on_render_done(0,1,old,old_gen)
    assert view._cache.get(0,1,0) is None
    view.unload_document(); doc.close()


def test_modified_later_page_invalidates_all_cached_pages(qapp, sample_pdf):
    doc=opened(sample_pdf); view=PDFView(); view.load_document(doc)
    view._cache.put(2,1,0,QPixmap(5,5))
    doc.pages.rotate([2],90)
    assert view._cache.get(2,1,0) is None
    view.unload_document(); doc.close()


def test_encrypted_view_renders_authenticated_memory(qapp, encrypted):
    doc=PDFDocument(); assert doc.open(str(encrypted),'user')
    view=PDFView(); view.load_document(doc)
    with patch('ui.viewer.pdf_view.PageRenderer.start_worker') as worker:
        view._render_visible_pages(); worker.assert_not_called()
    assert view._page_widgets[0]._pixmap is not None
    view.unload_document(); doc.close()


@pytest.mark.parametrize('rotation',[0,90,180,270])
def test_renderer_returns_thread_safe_image(qapp, rotation):
    from rendering.renderer import PageRenderer
    with fitz.open() as doc:
        page=doc.new_page(width=120,height=200)
        image=PageRenderer.render_page_image(page,1,rotation,2)
        assert isinstance(image,QImage) and image.devicePixelRatio()==2
        assert not image.isNull()


def test_replace_removes_original_content_and_preserves_vectors(sample_pdf,tmp_path):
    path=tmp_path/'replacement.pdf'
    with fitz.open() as src:
        p=src.new_page(width=320,height=200); p.insert_text((20,40),'REPLACEMENT')
        p.draw_rect(fitz.Rect(10,80,200,150),color=(1,0,0)); p.add_text_annot((50,60),'SOURCE NOTE'); src.save(path)
    doc=opened(sample_pdf); doc.pages.replace_pages([0],str(path))
    assert 'Sample page 1' not in doc.fitz_doc[0].get_text()
    assert 'REPLACEMENT' in doc.fitz_doc[0].get_text() and doc.fitz_doc[0].get_drawings()
    assert list(doc.fitz_doc[0].annots())
    assert doc.undo() and 'Sample page 1' in doc.fitz_doc[0].get_text()
    doc.close()


def test_flatten_preserves_appearance_and_removes_selected_annotations(sample_pdf):
    doc=opened(sample_pdf)
    p=doc.fitz_doc[0]; a=p.add_rect_annot(fitz.Rect(20,30,200,160)); a.set_colors(fill=(1,0,0)); a.update()
    doc.fitz_doc[1].add_text_annot((72,100),'Leave editable')
    before=doc.fitz_doc[0].get_pixmap().samples
    doc.annotations.flatten([0])
    assert not list(doc.fitz_doc[0].annots() or [])
    assert list(doc.fitz_doc[1].annots())
    assert before==doc.fitz_doc[0].get_pixmap().samples
    assert doc.undo() and list(doc.fitz_doc[0].annots())
    doc.close()


def test_duplicate_deletion_does_not_delete_extra_pages(sample_pdf):
    doc=opened(sample_pdf); doc.pages.delete([1,1])
    assert doc.page_count==2 and 'Sample page 3' in doc.fitz_doc[1].get_text()
    doc.close()


@pytest.mark.parametrize('order',[[0,0,1],[0,2],[-1,0,1]])
def test_invalid_reorder_does_not_modify_document(sample_pdf,order):
    doc=opened(sample_pdf)
    with pytest.raises(ValueError): doc.pages.reorder(order)
    assert doc.page_count==3 and not doc.is_modified and not doc.can_undo()
    doc.close()


def test_chinese_header_positions_margin_and_page_tokens(sample_pdf):
    doc=opened(sample_pdf)
    doc.pages.add_header_footer(header_left='左側',header='中央',header_right='右側 <<n>>',footer='第 {page} / {total} 頁',margin=30)
    page=doc.fitz_doc[0]; text=page.get_text()
    assert '左側' in text and '中央' in text and '右側 1' in text and '第 1 / 3 頁' in text
    blocks=page.get_text('blocks'); left=next(b for b in blocks if '左側' in b[4])
    assert abs(left[0]-30)<1
    doc.close()


def test_exports_do_not_silently_claim_pdfa(sample_pdf,tmp_path,monkeypatch):
    doc=opened(sample_pdf)
    path=tmp_path/'archive.pdf'
    monkeypatch.setattr('core.pdf_standards.find_executable', lambda kind: 'unused-ghostscript')
    monkeypatch.setattr('core.pdf_standards.validator_command', lambda: (_ for _ in ()).throw(RuntimeError('PDF/A validator missing')))
    with pytest.raises(RuntimeError,match='PDF/A'): doc.exports.export(str(path),'pdfa')
    assert not path.exists()
    doc.close()


def test_export_failure_preserves_destination(sample_pdf,tmp_path):
    doc=opened(sample_pdf); target=tmp_path/'existing.txt'; target.write_text('Original')
    def fail(path):
        open(path,'w').write('partial'); raise OSError('full')
    with patch.object(doc.exports,'export_txt',side_effect=fail):
        with pytest.raises(OSError): doc.exports.export(str(target),'txt')
    assert target.read_text()=='Original' and not list(tmp_path.glob('acropdf_export_*'))
    doc.close()


@pytest.mark.parametrize('fmt',['png','jpg','tiff'])
def test_images_use_chosen_name_and_return_actual_files(sample_pdf,tmp_path,fmt):
    doc=opened(sample_pdf); paths=doc.exports.export(str(tmp_path/('chosen.'+fmt)),fmt,dpi=72)
    assert len(paths)==3 and all(os.path.basename(p).startswith('chosen_p') for p in paths)
    assert all(os.path.isfile(p) for p in paths)
    with pytest.raises(FileExistsError): doc.exports.export(str(tmp_path/('chosen.'+fmt)),fmt,dpi=72)
    doc.close()


def test_image_render_failure_produces_no_partial_export(sample_pdf,tmp_path):
    doc=opened(sample_pdf)
    with patch.object(fitz.Page,'get_pixmap',side_effect=RuntimeError('render failure')):
        with pytest.raises(RuntimeError): doc.exports.export(str(tmp_path/'image.png'),'png')
    assert not list(tmp_path.glob('image_p*'))
    doc.close()


def test_excel_fallback_text_is_not_a_formula(tmp_path):
    import openpyxl
    path=tmp_path/'formula.pdf'
    with fitz.open() as d: d.new_page().insert_text((72,72),'=1+1'); d.save(path)
    doc=opened(path); out=tmp_path/'table.xlsx'; doc.exports.export(str(out),'xlsx')
    wb=openpyxl.load_workbook(out)
    assert wb.worksheets[0]['A1'].value=='=1+1' and wb.worksheets[0]['A1'].data_type=='s'
    doc.close()


def test_powerpoint_images_preserve_aspect_ratio(sample_pdf,tmp_path):
    from pptx import Presentation
    doc=opened(sample_pdf); out=tmp_path/'slides.pptx'; doc.exports.export(str(out),'pptx')
    deck=Presentation(out); shape=deck.slides[0].shapes[0]
    assert len(deck.slides)==3
    assert abs(shape.width/shape.height-doc.fitz_doc[0].rect.width/doc.fitz_doc[0].rect.height)<.01
    doc.close()


@pytest.mark.parametrize('endpoint,local',[('http://localhost:8080/v1',True),('http://127.0.0.1/v1',True),('http://[::1]/v1',True),('https://api.example.com/v1',False)])
def test_ai_endpoint_classification(endpoint,local):
    assert endpoint_is_local(endpoint)==local


@pytest.mark.parametrize('endpoint',['http://api.example.com/v1','https://user:secret@example.com/v1','file:///tmp/a','https://localhost.attacker.example:8080'])
def test_ai_does_not_confuse_remote_endpoints_with_local(endpoint):
    if endpoint.startswith('https://localhost.'):
        assert not endpoint_is_local(endpoint)
    else:
        with pytest.raises(ValueError): endpoint_is_local(endpoint)


def test_declining_remote_ai_does_not_send(window):
    with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):
        assert not confirm_document_transfer(window,'https://api.example.com/v1',3)


def test_accessibility_basic_check_never_claims_full_compliance(sample_pdf):
    from core.accessibility_engine import AccessibilityEngine
    doc=opened(sample_pdf)
    doc.fitz_doc.set_metadata({'title':'Title'})
    doc.fitz_doc.xref_set_key(doc.fitz_doc.pdf_catalog(),'Lang',fitz.get_pdf_str('zh-TW'))
    codes={issue['code'] for issue in AccessibilityEngine().validate_pdfua(doc.fitz_doc)}
    assert 'UA-002' not in codes and 'UA-004' in codes and 'UA-INCOMPLETE' in codes
    doc.close()


def test_pdfx_does_not_write_false_conformance(window,sample_pdf):
    from ui.dialogs.export.pdfx_dialog import PDFXDialog
    doc=opened(sample_pdf); before=doc.fitz_doc.tobytes(no_new_id=True)
    dialog=PDFXDialog(doc,window)
    with patch.object(QMessageBox,'warning') as info:
        dialog._export(); info.assert_called_once()
    assert dialog._worker is None
    assert doc.fitz_doc.tobytes(no_new_id=True)==before
    dialog.deleteLater(); doc.close()


def test_recovery_restores_edits_without_overwriting_source(sample_pdf,tmp_path):
    from core.recovery import RecoveryStore
    store=RecoveryStore(tmp_path/'recovery'); doc=opened(sample_pdf)
    original=sample_pdf.read_bytes(); doc.pages.rotate([0],90); store.capture(doc)
    entries=store.entries(); assert len(entries)==1
    restored=store.restore(entries[0])
    assert restored.is_modified and restored.path is None
    assert restored.fitz_doc[0].rotation==90 and sample_pdf.read_bytes()==original
    store.remove(restored); assert not store.entries()
    restored.close(); doc.close()


def test_recovery_retains_encryption_and_private_permissions(encrypted,tmp_path):
    from core.recovery import RecoveryStore
    store=RecoveryStore(tmp_path/'private'); doc=PDFDocument(); assert doc.open(str(encrypted),'user')
    doc.pages.rotate([0],90); store.capture(doc); entry=store.entries()[0]
    snapshot=store.directory/entry['file']
    if os.name != 'nt':
        assert snapshot.stat().st_mode & 0o777==0o600
    with fitz.open(snapshot) as check: assert check.needs_pass
    assert store.restore(entry) is None
    restored=store.restore(entry,password='user'); assert restored.fitz_doc[0].rotation==90
    restored.close();doc.close()


def test_recovery_detects_corruption_and_ignores_traversal(sample_pdf,tmp_path):
    import json
    from core.recovery import RecoveryStore
    store=RecoveryStore(tmp_path/'r');doc=opened(sample_pdf);doc.pages.rotate([0],90);store.capture(doc)
    entry=store.entries()[0];(store.directory/entry['file']).write_bytes(b'corrupt')
    with pytest.raises(ValueError,match='校驗'):store.restore(entry)
    entry['file']='../source.pdf';(store.directory/(entry['id']+'.json')).write_text(json.dumps(entry))
    assert not store.entries();doc.close()


def test_recovery_retains_only_latest_snapshot(sample_pdf,tmp_path):
    from core.recovery import RecoveryStore
    store=RecoveryStore(tmp_path/'r');doc=opened(sample_pdf);doc.pages.rotate([0],90);store.capture(doc)
    doc.pages.rotate([0],90);store.capture(doc)
    assert len(list(store.directory.glob('*.pdf')))==1
    restored=store.restore(store.entries()[0]);assert restored.fitz_doc[0].rotation==180
    restored.close();doc.close()


def test_save_removes_recovery_backup(window,sample_pdf):
    window.open_file(str(sample_pdf));window._current_doc().pages.rotate([0],90)
    window._capture_recovery();assert len(window._recovery.entries())==1
    window.save();assert not window._recovery.entries()


def test_render_budget_caps_extreme_page_scale(qapp):
    from rendering.renderer import PageRenderer
    with fitz.open() as doc:
        page=doc.new_page(width=10000,height=10000)
        pixels=fitz.Pixmap(fitz.csRGB,fitz.IRect(0,0,10,10),False)
        with patch.object(fitz.Page,'get_pixmap',return_value=pixels) as renderer:
            image=PageRenderer.render_page_image(page,8,0,2)
            matrix=renderer.call_args.kwargs['matrix']
            assert matrix.a*matrix.d*10000*10000<=24_000_001
            assert not image.isNull()


@pytest.mark.parametrize('theme',['light','dark'])
def test_text_and_primary_button_contrast(theme):
    def luminance(value):
        rgb=[int(value[i:i+2],16)/255 for i in (1,3,5)]
        linear=[v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4 for v in rgb]
        return sum(a*b for a,b in zip(linear,(.2126,.7152,.0722)))
    c=COLORS[theme]
    for foreground,background in [('text','bg'),('muted','bg'),('accent','bg'),('bg','accent')]:
        a,b=sorted((luminance(c[foreground]),luminance(c[background])))
        assert (b+.05)/(a+.05)>=4.5


def test_real_signing_is_atomic_and_signed_pdf_cannot_be_rewritten(sample_pdf,tmp_path):
    from datetime import datetime, timedelta, timezone
    from cryptography import x509
    from cryptography.x509.oid import NameOID
    from cryptography.hazmat.primitives import hashes,serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'AcroPDF test signer')])
    now=datetime.now(timezone.utc)
    cert=(x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
          .serial_number(x509.random_serial_number()).not_valid_before(now-timedelta(days=1))
          .not_valid_after(now+timedelta(days=1)).sign(key,hashes.SHA256()))
    credential=tmp_path/'test.p12'
    credential.write_bytes(pkcs12.serialize_key_and_certificates(b'test',key,cert,None,serialization.BestAvailableEncryption(b'secret')))
    doc=opened(sample_pdf);original=sample_pdf.read_bytes();signed=tmp_path/'signed.pdf'
    assert doc.signatures.sign_pdf(str(sample_pdf),str(signed),str(credential),b'secret')
    assert sample_pdf.read_bytes()==original
    check=opened(signed);before=signed.read_bytes()
    assert not check.save() and '簽章' in check.last_error
    assert signed.read_bytes()==before
    statuses=check.signatures.verify(str(signed))
    assert statuses[0]['valid'] and statuses[0]['intact']
    assert not statuses[0]['trusted'] and not statuses[0]['bottom_line']
    check.close();doc.close()


def test_signing_failures_do_not_truncate_destination(sample_pdf,tmp_path):
    doc=opened(sample_pdf);out=tmp_path/'existing.pdf';out.write_bytes(b'Original')
    assert not doc.signatures.sign_pdf(str(sample_pdf),str(out),str(tmp_path/'missing.pfx'),b'bad')
    assert out.read_bytes()==b'Original'
    original=sample_pdf.read_bytes()
    assert not doc.signatures.sign_pdf(str(sample_pdf),str(sample_pdf),'missing',b'bad')
    assert sample_pdf.read_bytes()==original
    doc.close()


@pytest.mark.parametrize('kind',['summary','translate'])
def test_running_ai_dialog_retains_worker_on_close(window,sample_pdf,kind):
    from unittest.mock import Mock
    from ui.dialogs.ai.summary_dialog import SummaryDialog
    from ui.dialogs.ai.translate_dialog import TranslateDialog
    doc=opened(sample_pdf);dialog=(SummaryDialog if kind=='summary' else TranslateDialog)(doc,window)
    worker=Mock();worker.isRunning.return_value=True;dialog._worker=worker
    event=QCloseEvent();dialog.closeEvent(event)
    assert not event.isAccepted();worker.requestInterruption.assert_called_once()
    dialog.reject();assert worker.requestInterruption.call_count==2
    dialog._worker=None;dialog.deleteLater();doc.close()


def test_redaction_removes_scanned_image_pixels_and_clears_history(tmp_path):
    path=tmp_path/'scan.pdf'
    with fitz.open() as d:
        p=d.new_page(width=200,height=200)
        pix=fitz.Pixmap(fitz.csRGB,fitz.IRect(0,0,100,100),False);pix.clear_with(120)
        p.insert_image(fitz.Rect(20,20,120,120),pixmap=pix)
        p.insert_text((20,150),'SECRET')
        d.save(path)
    doc=opened(path);rect=fitz.Rect(40,40,80,80)
    doc.annotations.add_redact(0,rect);doc.redaction.search_and_mark('SECRET')
    assert doc.can_undo() and doc.is_modified
    assert doc.redaction.apply_all()==2
    assert not doc.can_undo() and not doc.can_redo()
    assert doc.save(str(tmp_path/'clean.pdf'))
    with fitz.open(tmp_path/'clean.pdf') as check:
        assert 'SECRET' not in check[0].get_text()
        for info in check[0].get_images():
            embedded=fitz.Pixmap(check,info[0])
            # The corresponding source-image pixels must be replaced, not just covered by a rectangle.
            assert embedded.pixel(30,30)!=(120,120,120)
    doc.close()


def test_form_flatten_removes_fields_and_retains_value(sample_pdf):
    doc=opened(sample_pdf)
    doc.forms.add_field(0,fitz.Rect(72,100,260,135),'text','client','AcroPDF Form')
    before=doc.fitz_doc[0].get_pixmap().samples
    doc.forms.flatten_forms()
    assert not doc.forms.get_fields()
    assert 'AcroPDF Form' in doc.fitz_doc[0].get_text()
    assert before==doc.fitz_doc[0].get_pixmap().samples
    assert doc.undo() and doc.forms.get_fields()
    doc.close()


def test_pattern_mark_deduplicates_repeated_matches(sample_pdf):
    doc=opened(sample_pdf)
    doc.fitz_doc[0].insert_text((72,150),'0912345678 0912345678')
    results=doc.redaction.pattern_mark(r'09\d{8}')
    assert len(results)==2 and doc.redaction.get_redact_count()==2
    doc.close()


def test_fdf_roundtrip_handles_unicode_parentheses_and_backslashes(sample_pdf,tmp_path):
    doc=opened(sample_pdf)
    name='欄位 (姓名)';value='中文 (value) \\ path'
    doc.forms.add_field(0,fitz.Rect(72,100,260,135),'text',name,value)
    out=tmp_path/'fields.fdf';doc.forms.export_fdf(str(out))
    payload=out.read_bytes();assert payload.startswith(b'%FDF-1.2') and b'startxref' in payload
    doc.forms.fill_all({name:'different'});doc.forms.import_fdf(str(out))
    assert doc.forms.get_fields()[0]['field_value']==value
    doc.close()


def test_radio_groups_are_real_mutually_exclusive_fields(sample_pdf,tmp_path):
    doc=opened(sample_pdf)
    doc.forms.add_radio_button(0,fitz.Rect(50,100,70,120),'choice','alpha')
    doc.forms.add_radio_button(0,fitz.Rect(50,140,70,160),'choice','beta')
    widgets=list(doc.fitz_doc[0].widgets())
    assert all(w.field_type==fitz.PDF_WIDGET_TYPE_RADIOBUTTON for w in widgets)
    assert {w.field_name for w in widgets}=={'choice'}
    doc.forms.fill_all({'choice':'beta'});out=tmp_path/'radio.pdf';assert doc.save(str(out))
    with fitz.open(out) as check:
        widgets=list(check[0].widgets())
        states=[check.xref_get_key(w.xref,'AS')[1] for w in widgets]
        assert states==['/Off','/beta']
        parent=int(check.xref_get_key(widgets[0].xref,'Parent')[1].split()[0])
        assert check.xref_get_key(parent,'V')[1]=='/beta'
    doc.forms.fill_field(0,'choice','alpha')
    states=[doc.fitz_doc.xref_get_key(w.xref,'AS')[1] for w in doc.fitz_doc[0].widgets()]
    assert states==['/alpha','/Off']
    doc.close()


def test_text_edit_preserves_image_and_vector_background(sample_pdf):
    from core.content_editor import replace_text_span
    from ui.tools.text_edit_tool import _find_text_at
    doc = opened(sample_pdf)
    page = doc.fitz_doc[0]
    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 80), False)
    image.clear_with(180)
    page.insert_image(fitz.Rect(50, 50, 250, 130), pixmap=image)
    page.draw_rect(fitz.Rect(55, 55, 230, 85), color=(1, 0, 0))
    span = _find_text_at(page, fitz.Point(80, 68))
    replace_text_span(doc, 0, span, 'Short text')
    page = doc.fitz_doc[0]
    assert 'Short text' in page.get_text() and 'Sample page 1' not in page.get_text()
    assert page.get_images() and page.get_drawings()
    assert doc.undo() and 'Sample page 1' in doc.fitz_doc[0].get_text()
    doc.close()


def test_oversized_text_edit_keeps_original_and_history(sample_pdf):
    from core.content_editor import replace_text_span
    from ui.tools.text_edit_tool import _find_text_at
    doc = opened(sample_pdf)
    span = _find_text_at(doc.fitz_doc[0], fitz.Point(80, 68))
    before = doc.fitz_doc[0].get_text()
    with pytest.raises(ValueError, match='超出'):
        replace_text_span(doc, 0, span, 'Too long ' * 100)
    assert doc.fitz_doc[0].get_text() == before and not doc.can_undo() and not doc.is_modified
    doc.close()


def test_failed_edit_transaction_restores_encrypted_content_and_history(encrypted):
    doc = PDFDocument(); assert doc.open(str(encrypted), 'user')
    with pytest.raises(RuntimeError, match='injected failure'):
        with doc.edit_transaction('fault injection'):
            doc.fitz_doc[0].insert_text((72, 100), 'partial mutation')
            raise RuntimeError('injected failure')
    assert 'partial mutation' not in doc.fitz_doc[0].get_text()
    assert doc.is_encrypted and not doc.is_modified and not doc.can_undo() and not doc.can_redo()
    doc.close()


def test_image_deletion_preserves_overlay_text_and_vectors(sample_pdf):
    from core.content_editor import replace_image_region
    doc = opened(sample_pdf)
    page = doc.fitz_doc[0]
    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 80), False); image.clear_with(0)
    rect = fitz.Rect(50, 50, 250, 130)
    page.insert_image(rect, pixmap=image)
    page.draw_rect(rect, color=(1, 0, 0))
    replace_image_region(doc, 0, rect)
    page = doc.fitz_doc[0]
    assert 'Sample page 1' in page.get_text() and page.get_drawings()
    assert page.get_pixmap().pixel(200, 100) == (255, 255, 255)
    assert doc.undo() and doc.fitz_doc[0].get_images()
    doc.close()


def test_content_edit_does_not_consume_pending_redactions(sample_pdf):
    from core.content_editor import replace_image_region
    doc = opened(sample_pdf)
    doc.fitz_doc[0].add_redact_annot(fitz.Rect(70, 50, 180, 90))
    with pytest.raises(ValueError, match='待套用'):
        replace_image_region(doc, 0, fitz.Rect(200, 200, 300, 300))
    assert 'Sample page 1' in doc.fitz_doc[0].get_text()
    assert len(list(doc.fitz_doc[0].annots())) == 1
    assert not doc.can_undo()
    doc.close()


@pytest.mark.parametrize('rotation', [0, 90, 180, 270])
def test_rotated_overlay_geometry_roundtrips(qapp, rotation):
    with fitz.open() as doc:
        page = doc.new_page(); page.set_rotation(rotation)
        widget = PageWidget(0); widget._zoom = 2.0
        rect = fitz.Rect(50, 80, 180, 110)
        mapped = widget.pdf_rect_to_widget(rect, page)
        center = widget.widget_to_pdf(mapped.center(), page)
        assert abs(center.x - rect.tl.x - rect.width / 2) <= .5
        assert abs(center.y - rect.tl.y - rect.height / 2) <= .5
        widget.deleteLater()


def test_render_error_is_visible_and_retryable(qapp, monkeypatch):
    doc = PDFDocument(); doc.new()
    view = PDFView(); view.load_document(doc)
    from rendering.renderer import PageRenderer
    real_render = PageRenderer.render_page_sync
    def fail(*args): raise RuntimeError('injected rendering failure')
    monkeypatch.setattr(PageRenderer, 'render_page_sync', fail)
    view._render_visible_pages()
    widget = view._page_widgets[0]
    assert 'injected' in widget._render_error and 'injected' in widget.toolTip()
    monkeypatch.setattr(PageRenderer, 'render_page_sync', real_render)
    widget.render_retry.emit()
    assert widget._pixmap is not None and widget._render_error == ''
    view.unload_document(); view.close(); doc.close()


def test_base_font_is_not_misreported_as_embedded(sample_pdf):
    doc = opened(sample_pdf)
    assert any(f.name == 'Helvetica' and not f.embedded for f in doc.fonts.list_fonts())
    with pytest.raises(Exception): doc.fonts.embed_font('Helvetica', '/not-used.ttf')
    assert not doc.is_modified and not doc.can_undo()
    doc.close()


def test_accessibility_auto_tag_cannot_rewrite_existing_structure(qapp, sample_pdf, monkeypatch):
    from ui.dialogs.accessibility.accessibility_dialog import AccessibilityDialog
    doc = opened(sample_pdf)
    root = doc.fitz_doc.get_new_xref()
    doc.fitz_doc.update_object(root, '<< /Type /StructTreeRoot /K [] >>')
    doc.fitz_doc.xref_set_key(doc.fitz_doc.pdf_catalog(), 'StructTreeRoot', f'{root} 0 R')
    dialog = AccessibilityDialog(doc)
    assert not dialog._tag_btn.isEnabled()
    monkeypatch.setattr(QMessageBox, 'information', lambda *a: None)
    dialog._run_auto_tag()
    assert doc.fitz_doc.xref_get_key(doc.fitz_doc.pdf_catalog(), 'StructTreeRoot')[1] == f'{root} 0 R'
    assert not doc.can_undo()
    dialog.close(); doc.close()


@pytest.mark.parametrize('operation', ['fit_page', 'set_zoom', 'refresh', 'double_layout'])
def test_rebuild_preserves_last_page_and_visible_scroll(window, sample_pdf, qapp, operation):
    from app.constants import LayoutMode
    window.open_file(str(sample_pdf)); qapp.processEvents()
    view = window._current_view(); last = view._doc.page_count - 1
    view.go_to_page(last); qapp.processEvents()
    assert view.current_page() == last
    if operation == 'set_zoom': view.set_zoom(.9)
    elif operation == 'double_layout': view.set_layout_mode(LayoutMode.DOUBLE)
    else: getattr(view, operation)()
    qapp.processEvents(); qapp.processEvents()
    assert view.current_page() == last
    assert window._page_spin.value() == last + 1
    widget = next(page for page in view._page_widgets if page.page_num == last)
    top = widget.mapTo(view.viewport(), widget.rect().topLeft()).y()
    assert top < view.viewport().height() and top + widget.height() > 0


def test_pending_rebuild_restore_does_not_override_new_navigation(window, sample_pdf, qapp):
    window.open_file(str(sample_pdf)); qapp.processEvents()
    view = window._current_view(); view.go_to_page(1); qapp.processEvents()
    view.set_zoom(.9); view.go_to_page(view._doc.page_count - 1)
    qapp.processEvents(); qapp.processEvents()
    assert view.current_page() == view._doc.page_count - 1


@pytest.mark.skipif(os.name != 'nt', reason='Windows DACL verification')
def test_windows_recovery_acl_excludes_other_users(sample_pdf, tmp_path):
    import json, subprocess
    from core.recovery import RecoveryStore
    store = RecoveryStore(tmp_path / 'private-recovery')
    doc = opened(sample_pdf);doc.pages.rotate([0],90);store.capture(doc)
    entry = store.entries()[0]
    folder = str(store.directory).replace("'", "''")
    snapshot = str(store.directory / entry['file']).replace("'", "''")
    command = "$d=Get-Acl -LiteralPath '" + folder + "';$f=Get-Acl -LiteralPath '" + snapshot + "';@{protected=$d.AreAccessRulesProtected;sids=@($f.Access | ForEach-Object {$_.IdentityReference.Translate([System.Security.Principal.SecurityIdentifier]).Value})}|ConvertTo-Json -Compress"
    result = json.loads(subprocess.check_output(['powershell','-NoProfile','-Command',command],text=True))
    assert result['protected']
    assert not {'S-1-1-0', 'S-1-5-11', 'S-1-5-32-545'} & set(result['sids'])
    assert len(set(result['sids'])) == 3
    doc.close()

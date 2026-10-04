"""Output-based checks for content tagging, standards and Office/annotation exchange."""
from pathlib import Path
import io
import shutil
import fitz
import pikepdf
import pytest
from core.document import PDFDocument
from core.accessibility_engine import AccessibilityEngine
from core.pdf_standards import PDFStandards


def opened(path):
    doc=PDFDocument(); assert doc.open(str(path)); return doc


def test_tags_link_real_content_preserve_pixels_and_reorder(sample_pdf):
    doc=opened(sample_pdf);before=[p.get_pixmap().samples for p in doc.fitz_doc]
    engine=AccessibilityEngine(); assert engine.auto_tag(doc,'Structured document','en')==3
    tree=engine.get_structure_tree(doc.fitz_doc);assert len(tree.children[0].children)==3
    with pikepdf.open(io.BytesIO(doc._snapshot())) as pdf:
        assert pdf.Root.MarkInfo.Marked
        assert '/ParentTree' in pdf.Root.StructTreeRoot
        instructions=pikepdf.parse_content_stream(pdf.pages[0])
        assert any(str(op)=='BDC' and len(args)>1 and args[1].get('/MCID')==0 for args,op in instructions)
    assert [p.get_pixmap().samples for p in doc.fitz_doc]==before
    engine.reorder_structure(doc.fitz_doc,[2,0,1]);assert engine.get_structure_tree(doc.fitz_doc).children[0].children[0].page==2
    assert doc.undo() and engine.get_structure_tree(doc.fitz_doc) is None
    doc.close()


def test_tag_missing_alt_rolls_back_and_keeps_images(sample_pdf,tmp_path):
    from PIL import Image
    b=io.BytesIO();Image.new('RGB',(20,20),'red').save(b,format='PNG')
    doc=opened(sample_pdf);doc.fitz_doc[0].insert_image(fitz.Rect(72,150,200,220),stream=b.getvalue())
    before=doc.fitz_doc[0].get_pixmap().samples;engine=AccessibilityEngine()
    with pytest.raises(ValueError,match='替代文字'): engine.auto_tag(doc,'Image document')
    assert engine.get_structure_tree(doc.fitz_doc) is None and not doc.can_undo()
    assert doc.fitz_doc[0].get_pixmap().samples==before
    xref=doc.fitz_doc[0].get_images()[0][0];engine.auto_tag(doc,'Image document',image_alts={xref:'Red rectangle showing the selected color'})
    figure=engine.get_structure_tree(doc.fitz_doc).children[0].children[0].children[-1]
    assert figure.type=='Figure' and figure.alt_text.startswith('Red rectangle')
    doc.close()


def test_table_tags_keep_content_mcid_parenttree(sample_pdf):
    doc=opened(sample_pdf);engine=AccessibilityEngine();engine.auto_tag(doc,'Table document')
    table=engine.add_table_structure(doc.fitz_doc,0,fitz.Rect(50,40,300,100),1,1)
    root=engine.get_structure_tree(doc.fitz_doc)
    node=root.children[0].children[0].children[0]
    assert node.xref==table and node.type=='Table' and node.children[0].children[0].type=='TD'
    cell=node.children[0].children[0]
    engine.set_table_header(doc.fitz_doc,cell.xref,'Column')
    assert doc.fitz_doc.xref_get_key(cell.xref,'S')[1]=='/TH'
    assert doc.fitz_doc.xref_get_key(cell.xref,'A/Scope')[1]=='/Column'
    assert 'Sample page 1' in doc.fitz_doc[0].get_text()
    doc.close()


@pytest.mark.skipif(not shutil.which('java'),reason='Independent validator requires Java')
def test_pdfua_verified_real_fonts_and_content(tmp_path):
    fontfile=tmp_path/'NotoSans.ttf';fontfile.write_bytes(fitz.Font(fontname='notos').buffer)
    source=tmp_path/'source.pdf'
    with fitz.open() as f:
        p=f.new_page();p.insert_font(fontname='embedded',fontfile=str(fontfile));p.insert_text((72,72),'Accessible document',fontname='embedded');f.save(source)
    doc=opened(source);engine=AccessibilityEngine();engine.auto_tag(doc,'Accessible document','en')
    target=tmp_path/'accessible.pdf';report=PDFStandards(doc).export_pdfua(target)
    assert report['report']['jobs'][0]['validationResult'][0]['compliant']
    with fitz.open(target) as result: assert 'Accessible document' in result[0].get_text()
    doc.close()


@pytest.mark.skipif(not shutil.which('gs') or not Path('/System/Library/ColorSync/Profiles/Generic CMYK Profile.icc').exists(),reason='CMYK test profile and Ghostscript unavailable')
@pytest.mark.parametrize('level',['1','3','4'])
def test_pdfx_real_icc_and_fonts(sample_pdf,tmp_path,level):
    doc=opened(sample_pdf);target=tmp_path/'print.pdf'
    report=PDFStandards(doc).export_pdfx(target,'/System/Library/ColorSync/Profiles/Generic CMYK Profile.icc',level)
    assert report['page_count']==3
    with pikepdf.open(target) as result:
        assert result.Root.OutputIntents[0].DestOutputProfile.N==4
        assert result.Root.OutputIntents[0].DestOutputProfile.read_bytes()[16:20]==b'CMYK'
    doc.close()


def test_xfdf_chinese_note_and_text_highlight_roundtrip(sample_pdf,tmp_path):
    original=opened(sample_pdf);original.annotations.add_area_highlight(0,fitz.Rect(70,50,140,90))
    original.annotations.add_text_annot(0,fitz.Point(250,100),'中文審閱意見','審閱者')
    target=tmp_path/'review.xfdf';original.annotations.export_xfdf(target)
    new=opened(sample_pdf);assert new.annotations.import_xfdf(target)==2
    page=new.fitz_doc[0];annots=list(page.annots());assert annots[0].type[0]==fitz.PDF_ANNOT_HIGHLIGHT
    assert annots[1].info['content']=='中文審閱意見' and annots[1].info['title']=='審閱者'
    assert new.undo() and not list(new.fitz_doc[0].annots() or ())
    original.close();new.close()


def test_xfdf_failure_rolls_back_previous_entries(sample_pdf,tmp_path):
    file=tmp_path/'bad.xfdf';file.write_text('<xfdf xmlns="http://ns.adobe.com/xfdf/"><annots><text page="0" rect="20,20,40,40"><contents>First</contents></text><highlight page="0" rect="20,20,40,40" coords="1,2,3"/></annots></xfdf>')
    doc=opened(sample_pdf)
    with pytest.raises(ValueError): doc.annotations.import_xfdf(file)
    assert not doc.can_undo() and not list(doc.fitz_doc[0].annots() or ())
    doc.close()


def test_multiline_edit_keeps_other_content(sample_pdf):
    from core.content_editor import replace_text_span
    from ui.tools.text_edit_tool import _find_text_at
    doc=opened(sample_pdf);span=_find_text_at(doc.fitz_doc[0],fitz.Point(80,68))
    replace_text_span(doc,0,span,'First line\nSecond line')
    text=doc.fitz_doc[0].get_text();assert 'First line' in text and 'Second line' in text and 'Sample page 1' not in text
    assert 'Sample page 2' in doc.fitz_doc[1].get_text() and doc.undo()
    doc.close()


def test_multipage_tiff_keeps_all_images(tmp_path):
    from PIL import Image
    from core.file_converter import FileConverter
    source=tmp_path/'scan.tiff';Image.new('RGB',(100,120),'red').save(source,save_all=True,append_images=[Image.new('RGB',(80,140),'blue')])
    with FileConverter.convert(str(source)) as result:
        assert len(result)==2 and result[0].get_pixmap().pixel(5,5)[0]>200 and result[1].get_pixmap().pixel(5,5)[2]>200


def test_office_preserves_full_text_table_and_picture(tmp_path):
    from core.file_converter import FileConverter
    if not FileConverter._find_soffice(): pytest.skip('LibreOffice not installed')
    from docx import Document
    from PIL import Image
    word=Document();word.add_paragraph('Long content: '+'document word '*120);table=word.add_table(rows=1,cols=2);table.cell(0,0).text='Table column one';table.cell(0,1).text='Table column two'
    picture=tmp_path/'picture.png';Image.new('RGB',(120,80),'green').save(picture);word.add_picture(str(picture))
    source=tmp_path/'layout.docx';word.save(source)
    with FileConverter.convert(str(source)) as result:
        text=''.join(p.get_text() for p in result)
        assert text.count('document word')>=110 and 'Table column one' in text
        assert sum(len(p.get_images()) for p in result)>=1


def test_templates_write_user_data_and_keep_cjk(tmp_path,monkeypatch):
    from core.template_engine import TemplateEngine
    monkeypatch.setattr(TemplateEngine,'_TEMPLATE_DIR',tmp_path/'templates')
    engine=TemplateEngine()
    with fitz.open() as original:
        page=original.new_page();page.insert_text((72,72),'{name}');engine.save_as_template(original,'../中文範本')
    assert not (tmp_path.parent/'中文範本.pdf').exists()
    with engine.create_from_template('../中文範本',{'name':'王小明'}) as result:
        assert '王小明' in result[0].get_text() and '{name}' not in result[0].get_text()


def test_comparison_detects_different_page_sizes_and_added_pages(tmp_path):
    from core.compare_engine import CompareEngine
    a=tmp_path/'a.pdf';b=tmp_path/'b.pdf'
    with fitz.open() as doc:
        doc.new_page(width=200,height=300);doc.save(a)
    with fitz.open() as doc:
        doc.new_page(width=300,height=200);page=doc.new_page();page.insert_text((72,72),'New complete page');doc.save(b)
    engine=CompareEngine();result=engine.compare(str(a),str(b))
    assert result['pages_with_diff']==2 and result['extra_pages_in_b']==1
    assert result['page_diffs'][1].text_added==['New complete page']
    report=tmp_path/'report.pdf';engine.generate_diff_report(str(a),str(b),str(report),result)
    with fitz.open(report) as check: assert len(check)==2 and '此版本沒有此頁' in check[1].get_text()


def test_encrypted_page_extraction_preserves_passwords(tmp_path):
    source=tmp_path/'protected.pdf'
    with fitz.open() as doc:
        for index in range(2): doc.new_page().insert_text((72,72),f'Secret page {index+1}')
        doc.save(source,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='reader',owner_pw='owner')
    doc=PDFDocument();assert doc.open(str(source),'reader')
    target=tmp_path/'extracted.pdf';doc.pages.extract_pages([1],str(target))
    with fitz.open(target) as result:
        assert result.needs_pass and result.authenticate('reader') and 'Secret page 2' in result[0].get_text()
    outputs=doc.pages.split_by_range([(0,0),(1,1)],str(tmp_path/'split'))
    for path in outputs:
        with fitz.open(path) as result: assert result.needs_pass and result.authenticate('reader')
    doc.close()


@pytest.mark.skipif(__import__('sys').platform!='darwin',reason='Native Apple Vision requires macOS')
def test_native_ocr_chinese_preserves_glyphs_and_pixels(tmp_path):
    from acro_platform.mac import MacVisionOCR
    expected='商用文件品質審查 2026'
    with fitz.open() as text:
        page=text.new_page(width=600,height=120);page.insert_text((20,70),expected,fontname='china-t',fontsize=38)
        image=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).tobytes('png')
    with fitz.open() as scan:
        page=scan.new_page(width=600,height=120);page.insert_image(page.rect,stream=image);before=page.get_pixmap().samples
        MacVisionOCR().ocr_page(page,'chi_tra+eng',200)
        recognized=page.get_text().replace(' ','').replace('\n','')
        assert expected.replace(' ','')==recognized
        assert page.get_pixmap().samples==before


@pytest.mark.parametrize('name',['Approved','AsIs','Confidential','Departmental','Draft','Experimental','Expired','Final','ForComment','ForPublicRelease','NotApproved','NotForPublicRelease','Sold','TopSecret'])
def test_builtin_stamp_uses_correct_engine_code(sample_pdf,name):
    doc=opened(sample_pdf);page=doc.fitz_doc[0];annotation=doc.annotations.add_stamp(0,fitz.Rect(200,100,360,170),name)
    assert annotation.info['name']==name
    doc.close()


@pytest.mark.skipif(not shutil.which('java') or not shutil.which('gs'),reason='Standards engines unavailable')
def test_pdfa_complex_cjk_image_and_transparency(tmp_path):
    from PIL import Image
    image=Image.new('RGBA',(100,100),(255,50,10,120));b=io.BytesIO();image.save(b,format='PNG')
    source=tmp_path/'complex.pdf'
    with fitz.open() as f:
        p=f.new_page();p.insert_text((72,72),'商用文件品質審查',fontname='china-t',fontsize=18)
        p.insert_image(fitz.Rect(72,100,250,270),stream=b.getvalue());p.draw_rect(fitz.Rect(100,200,300,350),fill=(0,0,.8),fill_opacity=.4);f.save(source)
    doc=opened(source);target=tmp_path/'archived.pdf';report=PDFStandards(doc).export_pdfa(target)
    assert report['report']['jobs'][0]['validationResult'][0]['compliant']
    with fitz.open(target) as result:
        assert '商用文件品質審查' in result[0].get_text() and result[0].get_images() and result[0].get_pixmap().width>0
    doc.close()


def test_batch_bates_numbers_continue_across_files(sample_pdf,tmp_path):
    from core.batch_engine import BatchEngine
    second=tmp_path/'second.pdf';shutil.copy2(sample_pdf,second)
    out=tmp_path/'batch';result=BatchEngine().batch_process([str(sample_pdf),str(second)],'bates',{'output_dir':str(out),'prefix':'CASE-','start':10,'digits':3})
    assert all(item['ok'] for item in result)
    with fitz.open(result[0]['output']) as doc: assert 'CASE-012' in doc[2].get_text()
    with fitz.open(result[1]['output']) as doc: assert 'CASE-013' in doc[0].get_text()


def test_batch_same_basenames_both_have_distinct_outputs(sample_pdf,tmp_path):
    from core.batch_engine import BatchEngine
    other=tmp_path/'another';other.mkdir();copy=other/sample_pdf.name;shutil.copy2(sample_pdf,copy)
    result=BatchEngine().batch_process([str(sample_pdf),str(copy)],'watermark',{'output_dir':str(tmp_path/'result'),'text':'Reviewed'})
    assert all(item['ok'] for item in result) and result[0]['output']!=result[1]['output']
